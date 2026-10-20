#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5", "pypdf==6.21.1"]
# ///
"""Offline build of Wyoming county partisanship baselines and scenarios.

Reads only the retained official bytes listed in sources.json (each verified
against its SHA-256 before use).  No network access.  Writes only inside
senate/2026/wyoming/data/baselines/ and ../county-baselines.csv.

Run from any directory:
  uv run --with openpyxl==3.1.5 --with pypdf==6.21.1 python3 -I build.py
  ... build.py --print-tables   (also prints report-tables.txt, the Markdown
                                 tables used in county-partisanship.md;
                                 validate.py checks they appear verbatim)
"""
import argparse, collections, csv, datetime, gzip, hashlib, io, json, math, re, sys, warnings, zipfile
from pathlib import Path

import openpyxl, pypdf

warnings.filterwarnings('ignore', module='openpyxl')
ROOT = Path(__file__).resolve().parent
DATA = ROOT.parent
META = {m['id']: m for m in json.loads((ROOT / 'sources.json').read_text())}
PARSER_VERSION = 'wy-baselines-build/1'


def sha(b):
    return hashlib.sha256(b).hexdigest()


_RAW = {}


def raw(sid):
    if sid not in _RAW:
        m = META[sid]
        stored = (ROOT / m['path']).read_bytes()
        assert sha(stored) == m['stored_sha256'] and len(stored) == m['stored_bytes'], sid
        body = gzip.decompress(stored)
        assert sha(body) == m['original_sha256'] and len(body) == m['original_bytes'], sid
        _RAW[sid] = body
    return _RAW[sid]


def zmember(sid, name):
    return zipfile.ZipFile(io.BytesIO(raw(sid))).read(name)


def zinfo(sid):
    return [{'member': i.filename, 'bytes': i.file_size, 'zip_datetime': '%04d-%02d-%02dT%02d:%02d:%02d' % i.date_time,
             'sha256': sha(zipfile.ZipFile(io.BytesIO(raw(sid))).read(i.filename))}
            for i in zipfile.ZipFile(io.BytesIO(raw(sid))).infolist()]


def rows_of(xlsx_bytes, sheet):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    ws = wb[sheet]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def norm(s):
    return ' '.join(str(s).split()) if s is not None else ''


def as_int(v, where):
    if isinstance(v, bool) or v is None:
        raise AssertionError(f'non-numeric {v!r} at {where}')
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    raise AssertionError(f'non-integer {v!r} at {where}')


def colletter(c):
    s = ''
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def csvout(path, rows, fields=None):
    rows = list(rows)
    fields = fields or list(rows[0])
    with open(path, 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields, lineterminator='\n')
        w.writeheader()
        w.writerows(rows)


def dump(name, value):
    (ROOT / name).write_text(json.dumps(value, indent=2, ensure_ascii=False) + '\n')


# --------------------------------------------------------------------------
# Geography
# --------------------------------------------------------------------------
def counties():
    text = raw('census-wy-county-codes').decode()
    out = {}
    for r in csv.DictReader(io.StringIO(text), delimiter='|'):
        assert r['STATE'] == 'WY' and r['STATEFP'] == '56'
        out[r['COUNTYNAME'].removesuffix(' County')] = {'fips': r['STATEFP'] + r['COUNTYFP'], 'countyns': r['COUNTYNS'], 'census_name': r['COUNTYNAME']}
    assert len(out) == 23
    return out


COUNTY = counties()
NAMES = list(COUNTY)  # Census order == alphabetical == SOS order

# --------------------------------------------------------------------------
# Official XLSX summaries
# --------------------------------------------------------------------------
SUMMARY = {
    2020: ('sos-2020-general-zip', '2020_General_Statewide_Candidates_Summary.xlsx', 'Sheet1'),
    2022: ('sos-2022-general-zip', '2022_General_Summaries_Results.xlsx', 'Statewide Candidates'),
    2024: ('sos-2024-general-zip', '2024 General Results Summaries - OFFICIAL.xlsx', 'Statewide Candidates'),
}
BALLOTS = {
    2020: ('sos-2020-general-zip', '2020_General_Statewide_Total_Ballots_Cast_Summary.xlsx', 'Sheet1'),
    2022: ('sos-2022-general-zip', '2022_General_Summaries_Results.xlsx', 'Statewide Total Ballots Cast'),
    2024: ('sos-2024-general-zip', '2024 General Results Summaries - OFFICIAL.xlsx', 'Statewide Total Ballots Cast'),
}
SLUG = {'United States President': 'president', 'United States Senator': 'senate', 'United States Representative': 'house',
        'Governor': 'governor', 'Secretary of State': 'secretary-of-state', 'State Auditor': 'state-auditor',
        'State Treasurer': 'state-treasurer', 'Superintendent of Public Instruction': 'superintendent'}
PARTY = {'R': 'Republican', 'D': 'Democratic', 'L': 'Libertarian', 'I': 'Independent', 'C': 'Constitution', 'CONST': 'Constitution'}
SELECTED = {
    '2020-senate': {'role': 'primary same-seat baseline', 'seat': 'U.S. Senate Class II; the seat on the 2026 ballot (Lummis vs Ben David)'},
    '2024-senate': {'role': 'other-seat Senate baseline', 'seat': 'U.S. Senate Class I; the other Wyoming seat, not on the 2026 ballot (Barrasso vs Morrow)'},
    '2024-president': {'role': 'presidential comparison', 'seat': 'President; not a Senate contest'},
    '2022-governor': {'role': 'midterm statewide comparator', 'seat': 'Governor; 2022 midterm (no 2022 Wyoming U.S. Senate contest)'},
    '2022-house': {'role': 'supplementary midterm federal statewide contest; 2026 R Senate nominee on ballot', 'seat': 'U.S. House at-large (statewide); Hageman (R)'},
    '2024-house': {'role': 'supplementary federal statewide contest; 2026 R Senate nominee on ballot', 'seat': 'U.S. House at-large (statewide); Hageman (R)'},
}
HEADLINE = '2020-senate'


def header_kind(h):
    h = norm(h)
    if h == 'Write-Ins':
        return 'writein', '', ''
    if h == 'Overvotes':
        return 'over', '', ''
    if h == 'Undervotes':
        return 'under', '', ''
    m = re.match(r'^(.*?)\s*\(([^)]*)\)$', h)
    assert m, h
    code = m.group(2).strip().rstrip('.').upper()
    return 'candidate', m.group(1).strip(), code


def blocks_from(rows, header_r, title_r, start_col=1):
    """Return ordered blocks {title: [(col0, header)]} using forward-filled titles."""
    blocks = collections.OrderedDict()
    title = None
    for c in range(start_col, len(rows[header_r])):
        t = rows[title_r][c] if c < len(rows[title_r]) else None
        if isinstance(t, str) and norm(t):
            title = norm(t)
        h = rows[header_r][c]
        if h is None or not norm(h):
            continue
        assert title, (c, h)
        blocks.setdefault(title, []).append((c, h))
    return blocks


def check_block(title, cols):
    assert norm(cols[-1][1]) == 'Undervotes' and norm(cols[-2][1]) == 'Overvotes' and norm(cols[-3][1]) == 'Write-Ins', title


def find_header(rows, needle):
    for i, r in enumerate(rows):
        if any(isinstance(v, str) and norm(v) == needle for v in r):
            return i
    raise AssertionError(needle)


def title_row_above(rows, header_r):
    for i in range(header_r - 1, -1, -1):
        if any(isinstance(v, str) and norm(v) in SLUG for v in rows[i][1:]):
            return i
    raise AssertionError('no title row')


def parse_summary(year):
    sid, member, sheet = SUMMARY[year]
    rows = rows_of(zmember(sid, member), sheet)
    hr = find_header(rows, 'Write-Ins')
    tr = title_row_above(rows, hr)
    blocks = blocks_from(rows, hr, tr)
    for t, cols in blocks.items():
        check_block(t, cols)
    county_rows, total_row = {}, None
    for i in range(hr + 1, len(rows)):
        lab = norm(rows[i][0])
        if lab in COUNTY:
            assert lab not in county_rows, lab
            county_rows[lab] = i
        elif lab == 'Total':
            total_row = i
            break
    assert list(county_rows) == NAMES and total_row is not None, year
    cells = []
    for title, cols in blocks.items():
        key = f'{year}-{SLUG[title]}'
        for pos, (c, h) in enumerate(cols):
            kind, name, code = header_kind(h)
            for county, i in list(county_rows.items()) + [('Total', total_row)]:
                cells.append({'year': year, 'contest_key': key, 'contest_title': title, 'county': county,
                              'county_fips': COUNTY[county]['fips'] if county in COUNTY else '56',
                              'column_position': pos + 1, 'column_kind': kind, 'native_header': norm(h),
                              'candidate_name': name, 'party_code': code, 'party_label': PARTY.get(code, ''),
                              'votes': as_int(rows[i][c], f'{member}!{colletter(c + 1)}{i + 1}'),
                              'source_id': sid, 'zip_member': member, 'sheet': sheet, 'cell': f'{colletter(c + 1)}{i + 1}'})
    return cells


def parse_ballots(year):
    sid, member, sheet = BALLOTS[year]
    rows = rows_of(zmember(sid, member), sheet)
    out = {}
    for i, r in enumerate(rows):
        lab = norm(r[0])
        if lab in COUNTY or lab == 'Total':
            vals = [v for v in r[1:] if v is not None]
            assert len(vals) == 1, (year, lab, vals)
            out[lab] = {'ballots_cast': as_int(vals[0], f'{member}!{lab}'), 'cell': f'B{i + 1}', 'source_id': sid, 'zip_member': member, 'sheet': sheet}
    assert set(out) == set(NAMES) | {'Total'}
    return out


# --------------------------------------------------------------------------
# Precinct-by-precinct workbooks (independent county controls)
# --------------------------------------------------------------------------
def pbp_source(year, county):
    if year == 2020:
        return 'sos-2020-general-zip', f'2020_{county.replace(" ", "_")}_County_General_PbP.xlsx', 'Sheet1'
    if year == 2022:
        return 'sos-2022-general-zip', '2022_General_County_PbP.xlsx', county
    return 'sos-2024-general-zip', '2024 County General PbP - OFFICIAL.xlsx', county


_WB = {}


def pbp_rows(year, county):
    sid, member, sheet = pbp_source(year, county)
    if (sid, member) not in _WB:
        _WB[sid, member] = openpyxl.load_workbook(io.BytesIO(zmember(sid, member)), read_only=True, data_only=True)
    return [list(r) for r in _WB[sid, member][sheet].iter_rows(values_only=True)], (sid, member, sheet)


def label_repr(v):
    if isinstance(v, datetime.datetime):
        return 'excel-date:' + v.date().isoformat(), 'spreadsheet stores this precinct label as a date value'
    if isinstance(v, (int, float)):
        return str(v), 'numeric label'
    return norm(v), 'text label'


def parse_pbp(year, county, wanted):
    rows, (sid, member, sheet) = pbp_rows(year, county)
    hr = next(i for i, r in enumerate(rows) if norm(r[0]) == 'Precinct')
    tr = title_row_above(rows, hr)
    blocks = blocks_from(rows, hr, tr)
    datacols = [c for cols in blocks.values() for c, _ in cols]
    precincts, separators, total = [], [], None
    for i in range(hr + 1, len(rows)):
        lab = rows[i][0]
        if norm(lab) == 'Total':
            total = i
            break
        nums = [rows[i][c] for c in datacols if c < len(rows[i]) and rows[i][c] is not None]
        if not nums:
            if lab is not None:
                separators.append({'row': i + 1, 'label': norm(lab)})
            continue
        precincts.append(i)
    assert total is not None, (year, county)
    out = {}
    for title in wanted:
        cols = blocks[title]
        check_block(title, cols)
        col_out = []
        for pos, (c, h) in enumerate(cols):
            vals = [rows[i][c] for i in precincts]
            missing = sum(v is None for v in vals)
            # Non-numeric source text (e.g. a lone '.') is recorded and excluded,
            # never coerced to zero; the county comparison then shows whether the
            # published total is consistent with treating it as no votes.
            odd = [f'{colletter(c + 1)}{i + 1}={v!r}' for v, i in zip(vals, precincts) if isinstance(v, str)]
            psum = sum(as_int(v, f'{member}!{sheet}!{colletter(c + 1)}{i + 1}') for v, i in zip(vals, precincts) if v is not None and not isinstance(v, str))
            col_out.append({'position': pos + 1, 'header': norm(h), 'precinct_sum': psum, 'missing_cells': missing, 'nonnumeric_cells': '; '.join(odd),
                            'pbp_total': as_int(rows[total][c], f'{member}!{sheet}!{colletter(c + 1)}{total + 1}'),
                            'total_cell': f'{colletter(c + 1)}{total + 1}'})
        out[title] = col_out
    labels = [label_repr(rows[i][0]) for i in precincts]
    return {'source_id': sid, 'zip_member': member, 'sheet': sheet, 'precinct_rows': len(precincts),
            'unique_labels': len({l for l, _ in labels}), 'date_coerced_labels': sum(n.startswith('spreadsheet') for _, n in labels),
            'separator_rows': separators, 'blocks': out, 'labels': labels}


# --------------------------------------------------------------------------
# PDF cross-check of the same official summaries
# --------------------------------------------------------------------------
PDF = {2020: ('sos-2020-statewide-candidates-pdf', 'sos-2020-total-ballots-pdf'),
       2022: ('sos-2022-statewide-candidates-pdf', 'sos-2022-total-ballots-pdf'),
       2024: ('sos-2024-statewide-candidates-pdf', 'sos-2024-total-ballots-pdf')}
NUM = r'\d{1,3}(?:,\d{3})*'


def pdf_pages(sid):
    r = pypdf.PdfReader(io.BytesIO(raw(sid)))
    return [(p.extract_text() or '') for p in r.pages], {k: str(v) for k, v in (r.metadata or {}).items()}


def _tokens(line):
    """Numeric tokens with a deterministic text-order repair.

    pypdf sometimes emits the first column of a row as a separate text run
    glued onto the end of the line ("63,727 ... 6,860198,418").  A final
    token that starts exactly where the previous one ends is moved to the
    front, and the repair is reported."""
    ms = list(re.finditer(NUM, line))
    vals = [int(m.group(0).replace(',', '')) for m in ms]
    if len(ms) > 1 and ms[-1].start() == ms[-2].end():
        return [vals[-1]] + vals[:-1], 'reordered: glued final token is the first column'
    return vals, ''


def pdf_rows(text):
    rows = {}
    county_re = '|'.join(re.escape(n) for n in NAMES)
    lines = [l.strip() for l in text.splitlines()]
    for line in lines:
        m = re.match(rf'^({county_re})\s+((?:{NUM}\s*)+)$', line)
        if m:
            rows[m.group(1)] = (_tokens(m.group(2))[0], 'county line')
            continue
        m = re.match(rf'^Total\s+((?:{NUM}\s*)+)$', line)
        if m:
            rows['Total'] = (_tokens(m.group(1))[0], 'Total line')
            continue
        m = re.match(rf'^((?:{NUM}\s*)+)Total(?:\s+({NUM}))?$', line)
        if m:  # text-extraction order artifact: first total printed after (or no) label
            vals = _tokens(m.group(1))[0]
            first = [int(m.group(2).replace(',', ''))] if m.group(2) else []
            rows['Total'] = (first + vals, 'Total line reordered: extracted text places the first total after the label' if first else 'Total line: value printed before the label')
    if rows:
        return rows
    # Column layout: the 23 county names and Total appear as consecutive lines,
    # followed later by one numeric-only line per geography in the same order.
    seq = NAMES + ['Total']
    for i in range(len(lines) - len(seq) + 1):
        if lines[i:i + len(seq)] == seq:
            nums = [l for l in lines[i + len(seq):] if re.fullmatch(rf'(?:{NUM}\s*)+', l)]
            if len(nums) < len(seq):
                return {}
            for geo, l in zip(seq, nums):
                vals, note = _tokens(l)
                rows[geo] = (vals, 'column layout' + ('; ' + note if note else ''))
            return rows
    return {}


def pdf_crosscheck(cells, ballots):
    out = []
    by = collections.defaultdict(dict)
    for c in cells:
        by[c['contest_key'], c['county']][c['column_position']] = c['votes']
    titles = {c['contest_key']: c['contest_title'] for c in cells}
    for year, (cand_sid, bal_sid) in PDF.items():
        pages, meta = pdf_pages(cand_sid)
        for pno, text in enumerate(pages, 1):
            present = [t for t in SLUG if t in text and f'{year}-{SLUG[t]}' in titles]
            rows = pdf_rows(text)
            if not rows:
                out.append({'year': year, 'source_id': cand_sid, 'page': pno, 'contest_key': '', 'geography': '', 'status': 'no text table on page (cover or image)', 'pdf_values': '', 'xlsx_values': '', 'equal': ''})
                continue
            assert len(present) == 1, (year, pno, present)
            key = f'{year}-{SLUG[present[0]]}'
            for geo in NAMES + ['Total']:
                vals, note = rows[geo]
                x = [by[key, geo][p] for p in sorted(by[key, geo])]
                out.append({'year': year, 'source_id': cand_sid, 'page': pno, 'contest_key': key, 'geography': geo, 'status': note,
                            'pdf_values': ' '.join(map(str, vals)), 'xlsx_values': ' '.join(map(str, x)), 'equal': vals == x})
        pages, meta = pdf_pages(bal_sid)
        found = {}
        for pno, text in enumerate(pages, 1):
            for geo, (vals, note) in pdf_rows(text).items():
                found[geo] = (pno, vals, note)
        for geo in NAMES + ['Total']:
            if geo not in found:
                out.append({'year': year, 'source_id': bal_sid, 'page': '', 'contest_key': 'total-ballots-cast', 'geography': geo, 'status': 'not text-extractable in PDF', 'pdf_values': '', 'xlsx_values': ballots[year][geo]['ballots_cast'], 'equal': ''})
                continue
            pno, vals, note = found[geo]
            out.append({'year': year, 'source_id': bal_sid, 'page': pno, 'contest_key': 'total-ballots-cast', 'geography': geo, 'status': note,
                        'pdf_values': ' '.join(map(str, vals)), 'xlsx_values': ballots[year][geo]['ballots_cast'], 'equal': vals == [ballots[year][geo]['ballots_cast']]})
    return out


# --------------------------------------------------------------------------
# 2026 slate and rules
# --------------------------------------------------------------------------
def decode(b):
    for enc in ('utf-8-sig', 'cp1252'):
        try:
            return b.decode(enc), enc
        except UnicodeDecodeError:
            pass
    raise AssertionError('decode')


def strip_html(h):
    h = re.sub(r'<!--.*?-->', ' ', h, flags=re.S)
    h = re.sub(r'<(script|style)\b.*?</\1>', ' ', h, flags=re.S | re.I)
    return ' '.join(re.sub(r'<[^>]+>', ' ', h).replace('&amp;', '&').split())


def slate():
    text, enc = decode(raw('sos-2026-general-candidates-csv'))
    gen = list(csv.DictReader(io.StringIO(text)))
    senate = [r for r in gen if r['Office Sought'].strip() == 'UNITED STATES SENATOR']
    text, _ = decode(raw('sos-2026-primary-candidates-csv'))
    prim = list(csv.DictReader(io.StringIO(text)))
    psen = [r for r in prim if r['Office Sought'].strip().startswith('UNITED STATES SENATOR')]
    dup = collections.Counter(tuple(r.values()) for r in psen)
    pages, meta = pdf_pages('sos-2026-general-candidates-pdf')
    p1 = pages[0]
    seg = p1[p1.index('UNITED STATES SENATOR'):p1.index('UNITED STATES REPRESENTATIVE')]
    pdf_list = re.findall(r'^(Republican|Democratic|Libertarian|Constitution|Independent)\s+(.+?)\s+(?:P\.O\.|\d)', seg, flags=re.M)
    wpages, wmeta = pdf_pages('sos-2026-withdrawn-primary-pdf')
    wtext = '\n'.join(wpages)
    ind = strip_html(raw('sos-2026-independent-candidates-page').decode('utf-8', 'replace'))
    active = ind[ind.index('Active Independent Candidate Petitions'):ind.index('Office Information')]
    kd, kmeta = pdf_pages('sos-2026-key-dates-pdf')
    gdate = re.search(r'General Election\s*[.…\s]*\s*(November \d+, 2026)', kd[0]).group(1)
    law = '\n'.join(pdf_pages('wyoleg-title22-pdf')[0])
    law1 = ' '.join(law.split())

    def excerpt(start, end_marker):
        i = law1.index(start)
        j = law1.index(end_marker, i + len(start))
        return law1[i:j].strip()
    rules = {
        'W.S. 22-2-117(a)': excerpt('22-2-117. Vote required for election; ratification. (a)', ' As used in this subsection'),
        'W.S. 22-16-105': excerpt('22-16-105. Tie votes. A tie vote', ' 22-16-106.'),
        'W.S. 22-16-109(b)': excerpt('(b) There shall be a recount made of all the votes cast for any office in which', ' (c) There shall be a recount made of all the votes cast for any office if a losing'),
        'statute_source_id': 'wyoleg-title22-pdf', 'statute_last_modified': META['wyoleg-title22-pdf']['headers'].get('last-modified'),
    }
    # 2026 primary official results (U.S. Senator blocks)
    rows = rows_of(zmember('sos-2026-primary-zip', '2026 Primary Results Summaries - OFFICIAL.xlsx'), 'Statewide Candidates')
    hr = find_header(rows, 'Write-Ins')
    title, party, prim_rows = None, None, []
    tot = next(i for i in range(hr + 1, len(rows)) if norm(rows[i][0]) == 'Total')
    crow = {norm(rows[i][0]): i for i in range(hr + 1, tot) if norm(rows[i][0]) in COUNTY}
    assert list(crow) == NAMES
    for c in range(1, len(rows[hr])):
        if isinstance(rows[hr - 2][c], str) and norm(rows[hr - 2][c]):
            title = norm(rows[hr - 2][c]).replace(', Continued', '')
        if isinstance(rows[hr - 1][c], str) and norm(rows[hr - 1][c]):
            party = norm(rows[hr - 1][c])
        h = rows[hr][c]
        if title != 'United States Senator' or h is None:
            continue
        for geo, i in list(crow.items()) + [('Total', tot)]:
            prim_rows.append({'contest': '2026 primary United States Senator', 'party_primary': party, 'native_header': norm(h), 'county': geo,
                              'county_fips': COUNTY[geo]['fips'] if geo in COUNTY else '56', 'votes': as_int(rows[i][c], f'primary {colletter(c + 1)}{i + 1}'),
                              'source_id': 'sos-2026-primary-zip', 'zip_member': '2026 Primary Results Summaries - OFFICIAL.xlsx', 'sheet': 'Statewide Candidates', 'cell': f'{colletter(c + 1)}{i + 1}'})
    winners = {}
    for p in sorted({r['party_primary'] for r in prim_rows}):
        cand = [r for r in prim_rows if r['party_primary'] == p and r['county'] == 'Total' and norm(r['native_header']) not in ('Write-Ins', 'Overvotes', 'Undervotes')]
        for h in {r['native_header'] for r in prim_rows if r['party_primary'] == p}:
            assert sum(r['votes'] for r in prim_rows if r['party_primary'] == p and r['native_header'] == h and r['county'] != 'Total') == next(r['votes'] for r in prim_rows if r['party_primary'] == p and r['native_header'] == h and r['county'] == 'Total'), (p, h)
        cand.sort(key=lambda r: -r['votes'])
        winners[p] = {'winner_header': cand[0]['native_header'], 'winner_votes': cand[0]['votes'], 'runner_up_header': cand[1]['native_header'], 'runner_up_votes': cand[1]['votes'],
                      'candidates': {r['native_header']: r['votes'] for r in cand}}
    result = {
        'retrieved_utc': META['sos-2026-general-candidates-csv']['retrieval_utc'],
        'general_roster_csv': {'source_id': 'sos-2026-general-candidates-csv', 'url': META['sos-2026-general-candidates-csv']['requested_url'], 'last_modified': META['sos-2026-general-candidates-csv']['headers'].get('last-modified'), 'encoding': enc,
                               'rows_total': len(gen), 'us_senate_rows': [{k: r[k] for k in ('Office Sought', 'Party Affiliation', 'Ballot Name', 'Candidate Last Name', 'Candidate First Name', 'Date Filed', 'Date Withdrawn', 'Mailing City State & Zip')} for r in senate]},
        'general_roster_pdf': {'source_id': 'sos-2026-general-candidates-pdf', 'printed_timestamp': re.search(r'((?:Mon|Tues|Wednes|Thurs|Fri|Satur|Sun)day, \w+ \d+, 2026 - [\d:]+[AP]M)', p1).group(1), 'pdf_title': meta.get('/Title'), 'us_senate_entries': [list(x) for x in pdf_list]},
        'withdrawn_general_pdf': {'source_id': 'sos-2026-withdrawn-general-pdf', 'http_status': META['sos-2026-withdrawn-general-pdf']['http_status']},
        'withdrawn_primary_pdf': {'source_id': 'sos-2026-withdrawn-primary-pdf', 'mentions_us_senator': 'UNITED STATES SENATOR' in wtext, 'printed_timestamp': re.search(r'((?:Mon|Tues|Wednes|Thurs|Fri|Satur|Sun)day, \w+ \d+, 2026 - [\d:]+[AP]M)', wtext).group(1)},
        'independent_petitions_page': {'source_id': 'sos-2026-independent-candidates-page', 'active_section_text': active, 'mentions_senate': bool(re.search(r'senat', active, re.I))},
        'primary_roster_csv': {'source_id': 'sos-2026-primary-candidates-csv', 'last_modified': META['sos-2026-primary-candidates-csv']['headers'].get('last-modified'),
                               'us_senate_rows': [{k: r[k] for k in ('Office Sought', 'Party Affiliation', 'Ballot Name', 'Date Filed', 'Date Withdrawn')} for r in psen],
                               'exact_duplicate_rows': [{'ballot_name': k[14], 'count': v} for k, v in dup.items() if v > 1]},
        'primary_official_results': {'source_id': 'sos-2026-primary-zip', 'zip_last_modified': META['sos-2026-primary-zip']['headers'].get('last-modified'), 'winners': winners},
        'general_election_date': {'source_id': 'sos-2026-key-dates-pdf', 'value': gdate},
        'general_results_page_status': META['sos-2026-general-page']['http_status'],
        'rules': rules,
    }
    names = [(r['Party Affiliation'], r['Ballot Name']) for r in senate]
    assert names == [('REP', 'Harriet Hageman'), ('DEM', 'James Byrd')], names
    assert [tuple(x) for x in pdf_list] == [('Republican', 'Harriet Hageman'), ('Democratic', 'James Byrd')], pdf_list
    assert all(not r['Date Withdrawn'].strip() for r in senate)
    assert not result['withdrawn_primary_pdf']['mentions_us_senator'] and not result['independent_petitions_page']['mentions_senate']
    assert norm(winners['Republican']['winner_header']) == 'Harriet Hageman' and norm(winners['Democratic']['winner_header']) == 'James Byrd'
    assert gdate == 'November 3, 2026'
    return result, prim_rows


# --------------------------------------------------------------------------
# Scenario arithmetic
# --------------------------------------------------------------------------
RECOUNT_FACTOR = 1 / 199  # D-R >= 0.01*D  <=>  (D-R)/(D+R) >= 1/199 (two-candidate pair)
GOALS = [('tie', 0.0, 'D-minus-R pair margin = 0 (a mathematical tie; W.S. 22-16-105 decides ties by lot)'),
         ('D+1-pair-margin', 0.01, 'illustrative cushion: D-minus-R = 1% of D+R; needs +0.5 D-share pp beyond the tie'),
         ('D-outside-automatic-recount', RECOUNT_FACTOR, 'D-R >= 1% of D votes, the W.S. 22-16-109(b) automatic-recount boundary; (D-R)/(D+R) >= 1/199')]


def bounded_uniform(shares, weights, target_net):
    total = sum(weights)
    base = sum(t * (2 * s - 1) for t, s in zip(weights, shares))
    unb = (target_net - base) / (2 * total)
    if abs(target_net) > total:
        return {'status': 'impossible', 'unbounded_shift': unb, 'solved_shift': None, 'clipped': None, 'baseline_net': base}
    lo, hi = -1.0, 1.0
    for _ in range(200):
        mid = (lo + hi) / 2
        lead = sum(t * (2 * min(1, max(0, s + mid)) - 1) for t, s in zip(weights, shares))
        lo, hi = (mid, hi) if lead < target_net else (lo, mid)
    shift = (lo + hi) / 2
    clipped = sum(not 0 <= s + unb <= 1 for s in shares)
    if not clipped:
        shift = unb  # exact closed form when no county leaves [0,1]
    return {'status': 'bounded-feasible' if clipped else 'unbounded-feasible', 'unbounded_shift': unb, 'solved_shift': shift, 'clipped': clipped, 'baseline_net': base}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--print-tables', action='store_true')
    args = ap.parse_args()

    cells, ballots = [], {}
    for y in SUMMARY:
        cells += parse_summary(y)
        ballots[y] = parse_ballots(y)
    csvout(ROOT / 'county-candidates.csv', cells)

    # ---- per-contest county records ------------------------------------
    by = collections.defaultdict(list)
    for c in cells:
        by[c['contest_key'], c['county']].append(c)
    contests = sorted({c['contest_key'] for c in cells})
    state_controls = []
    for key in contests:
        cols = sorted({(c['column_position'], c['native_header'], c['column_kind']) for c in cells if c['contest_key'] == key})
        for pos, h, kind in cols:
            county_sum = sum(c['votes'] for n in NAMES for c in by[key, n] if c['column_position'] == pos)
            total = next(c['votes'] for c in by[key, 'Total'] if c['column_position'] == pos)
            state_controls.append({'contest_key': key, 'column_position': pos, 'native_header': h, 'column_kind': kind,
                                   'sum_of_23_county_rows': county_sum, 'published_total_row': total, 'difference': county_sum - total})
    for y in ballots:
        s = sum(ballots[y][n]['ballots_cast'] for n in NAMES)
        state_controls.append({'contest_key': f'{y}-total-ballots-cast', 'column_position': 1, 'native_header': 'Total Ballots Cast', 'column_kind': 'ballots',
                               'sum_of_23_county_rows': s, 'published_total_row': ballots[y]['Total']['ballots_cast'], 'difference': s - ballots[y]['Total']['ballots_cast']})

    def record(key, geo):
        cs = by[key, geo]
        year = int(key[:4])
        d = [c for c in cs if c['column_kind'] == 'candidate' and c['party_code'] == 'D']
        r = [c for c in cs if c['column_kind'] == 'candidate' and c['party_code'] == 'R']
        assert len(d) == 1 and len(r) == 1, key
        others = [c for c in cs if c['column_kind'] == 'candidate' and c['party_code'] not in ('D', 'R')]
        g = lambda k: sum(c['votes'] for c in cs if c['column_kind'] == k)
        D, R = d[0]['votes'], r[0]['votes']
        named_other = sum(c['votes'] for c in others)
        wi, over, under = g('writein'), g('over'), g('under')
        cv = D + R + named_other + wi
        bal = ballots[year][geo]['ballots_cast']
        pair = D + R
        return {'contest_key': key, 'year': year, 'contest_title': cs[0]['contest_title'], 'baseline_role': SELECTED.get(key, {}).get('role', 'context only'),
                'seat_or_office_note': SELECTED.get(key, {}).get('seat', ''), 'county': geo, 'county_fips': COUNTY[geo]['fips'] if geo in COUNTY else '56',
                'd_candidate': d[0]['candidate_name'], 'r_candidate': r[0]['candidate_name'],
                'named_other_candidates': '; '.join(f"{c['candidate_name']} ({c['party_code']})" for c in others),
                'd_votes': D, 'r_votes': R, 'named_other_votes': named_other, 'writein_votes': wi,
                'contest_votes_candidates_plus_writein': cv, 'over_votes': over, 'under_votes': under,
                'contest_votes_plus_over_under': cv + over + under, 'ballots_cast': bal, 'ballots_identity_holds': cv + over + under == bal,
                'pair_votes': pair, 'pair_participation_of_ballots': pair / bal, 'd_pair_share': D / pair, 'd_minus_r_net': D - R,
                'd_minus_r_pair_margin': (D - R) / pair, 'd_minus_r_allcandidate_margin': (D - R) / cv,
                'nonmajor_share_of_contest_votes': (named_other + wi) / cv, 'pair_winner': 'D' if D > R else 'R' if R > D else 'tie'}

    county_contests = []
    for key in SELECTED:
        recs = [record(key, n) for n in NAMES]
        st = record(key, 'Total')
        tp = sum(r['pair_votes'] for r in recs)
        tn = sum(r['d_minus_r_net'] for r in recs)
        assert tp == st['pair_votes'] and tn == st['d_minus_r_net']
        for r in recs:
            r['pair_weight'] = r['pair_votes'] / tp
            r['share_of_statewide_net'] = r['d_minus_r_net'] / tn
            a = abs(r['d_minus_r_pair_margin']) * 100
            r['closeness_band'] = 'within 5 margin pp' if a <= 5 else 'over 5 to 10 margin pp' if a <= 10 else 'over 10 margin pp'
            r['net_change_per_plus10pct_pair_turnout'] = 0.1 * r['pair_votes'] * (2 * r['d_pair_share'] - 1)
        st.update({'pair_weight': 1.0, 'share_of_statewide_net': 1.0, 'closeness_band': '', 'net_change_per_plus10pct_pair_turnout': 0.1 * st['d_minus_r_net']})
        county_contests += recs + [st]
    csvout(ROOT / 'county-contests.csv', county_contests)
    dump('candidate-headers.json', {key: sorted({(c['column_position'], c['native_header'], c['column_kind'], c['party_code']) for c in cells if c['contest_key'] == key}) for key in contests})

    # ---- precinct controls ---------------------------------------------
    pchecks, plabels, pcounts = [], [], []
    sel_titles = collections.defaultdict(list)
    for key in SELECTED:
        y = int(key[:4])
        sel_titles[y].append(next(c['contest_title'] for c in cells if c['contest_key'] == key))
    for y in SUMMARY:
        for n in NAMES:
            p = parse_pbp(y, n, sel_titles[y])
            pcounts.append({'year': y, 'county': n, 'county_fips': COUNTY[n]['fips'], 'source_id': p['source_id'], 'zip_member': p['zip_member'], 'sheet': p['sheet'],
                            'precinct_rows': p['precinct_rows'], 'unique_precinct_labels': p['unique_labels'], 'date_coerced_labels': p['date_coerced_labels'],
                            'separator_rows': '; '.join(f"row {s['row']}: {s['label']}" for s in p['separator_rows'])})
            for lab, note in p['labels']:
                plabels.append({'year': y, 'county': n, 'county_fips': COUNTY[n]['fips'], 'native_precinct_label': lab, 'label_note': note})
            for title, cols in p['blocks'].items():
                key = f'{y}-{SLUG[title]}'
                summ = sorted(by[key, n], key=lambda c: c['column_position'])
                assert [norm(c['native_header']) for c in summ] == [c['header'] for c in cols], (key, n)
                for c, s in zip(cols, summ):
                    pchecks.append({'contest_key': key, 'county': n, 'county_fips': COUNTY[n]['fips'], 'column_position': c['position'], 'native_header': c['header'],
                                    'precinct_row_sum': c['precinct_sum'], 'precinct_missing_cells': c['missing_cells'], 'precinct_nonnumeric_cells': c['nonnumeric_cells'], 'pbp_total_row': c['pbp_total'],
                                    'county_summary_value': s['votes'], 'precinct_sum_equals_summary': c['precinct_sum'] == s['votes'],
                                    'pbp_total_equals_summary': c['pbp_total'] == s['votes'], 'pbp_source': f"{p['zip_member']}!{p['sheet']}!{c['total_cell']}"})
    csvout(ROOT / 'precinct-aggregation-checks.csv', pchecks)
    csvout(ROOT / 'precinct-counts.csv', pcounts)
    csvout(ROOT / 'precinct-labels.csv', plabels)

    # ---- PDF cross-check --------------------------------------------------
    pdfc = pdf_crosscheck(cells, ballots)
    csvout(ROOT / 'pdf-crosscheck.csv', pdfc)

    # ---- 2026 slate --------------------------------------------------------
    sl, prim_rows = slate()
    dump('slate-2026.json', sl)
    csvout(ROOT / 'primary-2026-us-senate.csv', prim_rows)

    # ---- scenarios ---------------------------------------------------------
    cc = {(r['contest_key'], r['county']): r for r in county_contests}
    mixes = [('native-pair', None)] + [(f'ballots-{y}', y) for y in (2020, 2022, 2024)]
    settings, state_rows, county_rows, sel_rows = [], [], [], []
    for pref in SELECTED:
        py = int(pref[:4])
        shares = [cc[pref, n]['d_pair_share'] for n in NAMES]
        for mix, my in mixes:
            if my == py:
                continue  # identical to native-pair: same ballots, same participation
            if my is None:
                T = [float(cc[pref, n]['pair_votes']) for n in NAMES]
                rule = 'exact historical D+R votes of the preference contest'
            else:
                T = [ballots[my][n]['ballots_cast'] * cc[pref, n]['pair_votes'] / cc[pref, n]['ballots_cast'] for n in NAMES]
                rule = f'{my} county total ballots cast x preference-contest county (D+R)/ballots cast'
            sid = f'{pref}__{mix}'
            total = sum(T)
            settings.append({'scenario_id': sid, 'preference_contest': pref, 'volume_mix': mix, 'pair_pool_rule': rule,
                             'status': 'historical composition sensitivity, not a 2026 turnout projection'})
            for goal, m, desc in GOALS:
                target = m * total
                res = bounded_uniform(shares, T, target)
                closed = (target - res['baseline_net']) / (2 * total)
                state_rows.append({'scenario_id': sid, 'preference_contest': pref, 'volume_mix': mix, 'goal': goal, 'goal_definition': desc,
                                   'pair_pool': total, 'baseline_d_minus_r_net': res['baseline_net'], 'baseline_pair_margin': res['baseline_net'] / total,
                                   'target_net': target, 'unbounded_d_share_shift': res['unbounded_shift'], 'solved_d_share_shift': res['solved_shift'],
                                   'closed_form_shift': closed, 'clipped_counties': res['clipped'], 'status': res['status'],
                                   'margin_point_change_pp': 200 * res['solved_shift'], 'net_vote_change_needed': target - res['baseline_net']})
                for n, s, t in zip(NAMES, shares, T):
                    ts = min(1, max(0, s + res['solved_shift']))
                    county_rows.append({'scenario_id': sid, 'goal': goal, 'county': n, 'county_fips': COUNTY[n]['fips'], 'baseline_d_pair_share': s,
                                        'scenario_pair_votes': t, 'pair_weight': t / total, 'baseline_net': t * (2 * s - 1), 'target_d_pair_share': ts,
                                        'target_pair_margin': 2 * ts - 1, 'target_net': t * (2 * ts - 1), 'required_d_share_change': ts - s,
                                        'required_margin_change': 2 * (ts - s), 'net_gain_per_plus1pp_d_share': 0.02 * t})
                    # selected county alone; every other county at baseline
                    fixed = sum(tt * (2 * ss - 1) for nn, ss, tt in zip(NAMES, shares, T) if nn != n)
                    q = (target - fixed + t) / (2 * t)
                    sel_rows.append({'scenario_id': sid, 'goal': goal, 'county': n, 'county_fips': COUNTY[n]['fips'], 'selected_pair_votes': t,
                                     'fixed_other_counties_net': fixed, 'target_net': target, 'baseline_d_pair_share': s, 'required_d_pair_share': q,
                                     'required_change_pp': 100 * (q - s), 'status': 'already-secured' if q < 0 else 'impossible' if q > 1 else 'feasible',
                                     'max_net_if_county_100pct_d': fixed + t})
    dump('scenario-settings.json', {'parser_version': PARSER_VERSION, 'goals': [{'goal': g, 'pair_margin': m, 'definition': d} for g, m, d in GOALS],
                                    'scenarios': settings, 'notes': [
                                        'Pair = D+R of the preference contest. Shares are fractions; *_pp columns are percentage points.',
                                        'Uniform additive D-share shift; bounded solver clips to [0,1] and reports clipped county counts (none clip in these baselines).',
                                        'Mix rows rescale county pair pools with another year\'s county ballots cast; mixes equal to the preference year are omitted because they reproduce native-pair exactly.',
                                        'No 2026 turnout projection, probability, or live remaining-vote estimate is implied.']})
    csvout(ROOT / 'state-scenarios.csv', state_rows)
    csvout(ROOT / 'county-target-scenarios.csv', county_rows)
    csvout(ROOT / 'selected-county-break-even.csv', sel_rows)

    # ---- turnout sensitivity (headline and other Senate baselines) -------
    sens = []
    for pref in ('2020-senate', '2024-senate', '2024-president'):
        recs = [cc[pref, n] for n in NAMES]
        T = [r['pair_votes'] for r in recs]
        sh = [r['d_pair_share'] for r in recs]
        net = sum(r['d_minus_r_net'] for r in recs)
        tot = sum(T)
        base_tie = bounded_uniform(sh, [float(t) for t in T], 0.0)['solved_shift']
        scaled = {f: bounded_uniform(sh, [f * t for t in T], 0.0)['solved_shift'] for f in (0.6, 0.75, 1.25)}
        dled = [n for n, r in zip(NAMES, recs) if r['d_pair_share'] > 0.5]
        dled_margin = {n: 2 * cc[pref, n]['d_pair_share'] - 1 for n in dled}
        # extra pair votes in all D-led counties (at their own baseline shares,
        # scaled proportionally to their baseline pair votes) needed to tie
        dnet_per_vote = sum(cc[pref, n]['pair_votes'] * dled_margin[n] for n in dled) / sum(cc[pref, n]['pair_votes'] for n in dled)
        extra = -net / dnet_per_vote
        mix_shifts = {r['volume_mix']: r['solved_d_share_shift'] for r in state_rows if r['preference_contest'] == pref and r['goal'] == 'tie'}
        sens.append({'preference_contest': pref, 'pair_pool': tot, 'baseline_net': net, 'net_votes_per_1pp_statewide_d_share': 0.02 * tot,
                     'tie_shift_native': base_tie, 'tie_shift_uniformly_scaled_turnout': scaled,
                     'uniform_plus10pct_turnout_net_change': 0.1 * net,
                     'd_led_counties': dled, 'd_led_counties_baseline_net': sum(cc[pref, n]['d_minus_r_net'] for n in dled),
                     'd_led_counties_baseline_pair_votes': sum(cc[pref, n]['pair_votes'] for n in dled),
                     'extra_pair_votes_in_d_led_counties_to_tie_at_baseline_shares': extra,
                     'extra_as_multiple_of_d_led_baseline_pair_votes': extra / sum(cc[pref, n]['pair_votes'] for n in dled),
                     'tie_shift_by_volume_mix': mix_shifts,
                     'tie_shift_range_across_mixes_pp': 100 * (max(mix_shifts.values()) - min(mix_shifts.values()))})
    dump('turnout-sensitivity.json', {'notes': ['Preferences fixed at the named contest. Uniform scaling of every county pool leaves the share threshold unchanged; only composition (relative county weights) moves it.',
                                                'Extra D-led-county votes are added at those counties\' own baseline two-party shares; this isolates turnout from persuasion and is not a projection.'],
                                      'rows': sens})

    # ---- independent priority ranking (2024 Senate contest votes) --------
    rk = sorted(NAMES, key=lambda n: (-cc['2024-senate', n]['contest_votes_candidates_plus_writein'], COUNTY[n]['fips']))
    den = sum(cc['2024-senate', n]['contest_votes_candidates_plus_writein'] for n in NAMES)
    cum, rank, k80 = 0, {}, None
    for i, n in enumerate(rk, 1):
        cum += cc['2024-senate', n]['contest_votes_candidates_plus_writein']
        rank[n] = (i, cum)
        if k80 is None and cum * 5 >= den * 4:
            k80 = i
    dump('priority-check.json', {'basis': '2024 U.S. Senate candidate+write-in votes (excludes over/under)', 'denominator': den, 'minimum_count_reaching_80pct': k80,
                                 'counties_in_80pct_set': rk[:k80], 'top15_votes': sum(cc['2024-senate', n]['contest_votes_candidates_plus_writein'] for n in rk[:15]),
                                 'status': 'independent partisanship-worker computation for reconciliation; the turnout worker owns the accepted ranking'})

    # ---- wide county CSV -----------------------------------------------------
    hl_native = {(r['county']): r for r in county_rows if r['scenario_id'] == f'{HEADLINE}__native-pair'}
    tie = {r['county']: r for r in county_rows if r['scenario_id'] == f'{HEADLINE}__native-pair' and r['goal'] == 'tie'}
    d1 = {r['county']: r for r in county_rows if r['scenario_id'] == f'{HEADLINE}__native-pair' and r['goal'] == 'D+1-pair-margin'}
    t24 = {r['county']: r for r in county_rows if r['scenario_id'] == '2024-senate__native-pair' and r['goal'] == 'D+1-pair-margin'}
    t22 = {r['county']: r for r in county_rows if r['scenario_id'] == f'{HEADLINE}__ballots-2022' and r['goal'] == 'D+1-pair-margin'}
    wide = []
    for n in NAMES:
        row = {'county': n, 'county_fips': COUNTY[n]['fips'], 'census_countyns': COUNTY[n]['countyns'], 'sos_county_label': n,
               'rank_2024_senate_contest_votes': rank[n][0], 'in_min_80pct_set_2024_senate': rank[n][0] <= k80}
        for key in SELECTED:
            p = key.replace('-', '_') + '_'
            r = cc[key, n]
            for f in ('d_votes', 'r_votes', 'named_other_votes', 'writein_votes', 'contest_votes_candidates_plus_writein', 'over_votes', 'under_votes', 'ballots_cast',
                      'pair_votes', 'pair_participation_of_ballots', 'd_pair_share', 'd_minus_r_net', 'd_minus_r_pair_margin', 'd_minus_r_allcandidate_margin',
                      'pair_winner', 'closeness_band', 'pair_weight', 'share_of_statewide_net'):
                row[p + f] = r[f]
        row.update({'scn_2020senate_native_tie_target_d_share': tie[n]['target_d_pair_share'], 'scn_2020senate_native_dplus1_target_d_share': d1[n]['target_d_pair_share'],
                    'scn_2020senate_native_dplus1_target_net': d1[n]['target_net'], 'scn_2020senate_native_net_gain_per_plus1pp': d1[n]['net_gain_per_plus1pp_d_share'],
                    'scn_2020senate_net_change_per_plus10pct_pair_turnout': cc[HEADLINE, n]['net_change_per_plus10pct_pair_turnout'],
                    'scn_2024senate_native_dplus1_target_d_share': t24[n]['target_d_pair_share'], 'scn_2020senate_2022mix_pair_votes': t22[n]['scenario_pair_votes'],
                    'scn_2020senate_2022mix_dplus1_target_d_share': t22[n]['target_d_pair_share'],
                    'source_ids': 'sos-2020-general-zip; sos-2022-general-zip; sos-2024-general-zip; census-wy-county-codes', 'research_date': '2026-10-18'})
        wide.append(row)
    csvout(DATA / 'county-baselines.csv', wide)

    dump('zip-members.json', {sid: zinfo(sid) for sid in ('sos-2020-general-zip', 'sos-2022-general-zip', 'sos-2024-general-zip', 'sos-2026-primary-zip')})
    summary = {'parser_version': PARSER_VERSION, 'state_controls_nonzero': sum(r['difference'] != 0 for r in state_controls),
               'precinct_cells': len(pchecks), 'precinct_mismatches': sum(not (r['precinct_sum_equals_summary'] and r['pbp_total_equals_summary']) for r in pchecks),
               'pdf_rows_compared': sum(r['equal'] in (True, False) for r in pdfc), 'pdf_mismatches': sum(r['equal'] is False for r in pdfc),
               'pdf_rows_not_text': sum(r['equal'] == '' for r in pdfc)}
    csvout(ROOT / 'state-controls.csv', state_controls)
    dump('build-summary.json', summary)
    print(json.dumps(summary))
    tables = render_tables(county_contests, state_rows, sel_rows, sl, wide, k80, rk, den, sens)
    (ROOT / 'report-tables.txt').write_text(tables + '\n')
    if args.print_tables:
        print(tables)


def f0(x):
    return f'{x:,.0f}'


def fs(x, d=2):
    return f'{x:+.{d}f}'


def render_tables(county_contests, state_rows, sel_rows, sl, wide, k80, rk, den, sens):
    cc = {(r['contest_key'], r['county']): r for r in county_contests}
    out = []
    out.append('<!-- table:slate -->')
    out.append('| 2026 U.S. Senate ballot | Party (roster) | Date filed | 2026 primary votes | Primary runner-up | Source |')
    out.append('| --- | --- | --- | ---: | --- | --- |')
    w = sl['primary_official_results']['winners']
    for r in sl['general_roster_csv']['us_senate_rows']:
        pw = w['Republican' if r['Party Affiliation'] == 'REP' else 'Democratic']
        out.append(f"| {r['Ballot Name']} | {r['Party Affiliation']} | {r['Date Filed']} | {pw['winner_votes']:,} | {norm(pw['runner_up_header'])} ({pw['runner_up_votes']:,}) | general roster CSV/PDF; primary official XLSX |")
    out.append('')
    out.append('<!-- table:turnout -->')
    out.append('| Preference baseline | Net votes per 1 D-share pp | Tie shift pp, actual pools | Tie shift pp, all pools ×0.60 / ×1.25 | Uniform +10% turnout net change | D-led counties | Extra D-led pair votes to tie | Multiple of their baseline pair votes | Tie-shift range across historical mixes pp |')
    out.append('| --- | ---: | ---: | ---: | ---: | --- | ---: | ---: | ---: |')
    for x in sens:
        out.append(f"| {x['preference_contest']} | {x['net_votes_per_1pp_statewide_d_share']:,.2f} | {100 * x['tie_shift_native']:+.6f} | {100 * x['tie_shift_uniformly_scaled_turnout'][0.6]:+.6f} / {100 * x['tie_shift_uniformly_scaled_turnout'][1.25]:+.6f} | {x['uniform_plus10pct_turnout_net_change']:+,.1f} | {', '.join(x['d_led_counties'])} | {x['extra_pair_votes_in_d_led_counties_to_tie_at_baseline_shares']:,.0f} | {x['extra_as_multiple_of_d_led_baseline_pair_votes']:.1f} | {x['tie_shift_range_across_mixes_pp']:.6f} |")
    out.append('')
    out.append('<!-- table:priority -->')
    acc = 0
    out.append('| Rank | County | 2024 Senate candidate+write-in votes | Cumulative % |')
    out.append('| ---: | --- | ---: | ---: |')
    for i, n in enumerate(rk[:k80], 1):
        v = cc['2024-senate', n]['contest_votes_candidates_plus_writein']
        acc += v
        out.append(f'| {i} | {n} | {v:,} | {100 * acc / den:.4f} |')
    out.append('')
    out.append('<!-- table:statewide -->')
    out.append('| Baseline | D candidate | R candidate | D | R | Named other | Write-in | Over | Under | Ballots cast | D two-party % | D−R pair margin pp | D−R net | Counties D led |')
    out.append('| --- | --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: |')
    for key in SELECTED:
        s = cc[key, 'Total']
        dl = [n for n in NAMES if cc[key, n]['pair_winner'] == 'D']
        out.append(f"| {key} | {s['d_candidate']} | {s['r_candidate']} | {f0(s['d_votes'])} | {f0(s['r_votes'])} | {f0(s['named_other_votes'])} | {f0(s['writein_votes'])} | {f0(s['over_votes'])} | {f0(s['under_votes'])} | {f0(s['ballots_cast'])} | {100 * s['d_pair_share']:.4f} | {100 * s['d_minus_r_pair_margin']:+.4f} | {s['d_minus_r_net']:+,} | {', '.join(dl) or 'none'} |")
    out.append('')
    out.append('<!-- table:scenarios -->')
    out.append('| Preference baseline | County volume mix | Pair pool | Baseline D−R net | Tie: D-share shift pp | D+1: D-share shift pp | Recount-free D win: D-share shift pp | D+1 net-vote gain needed | Clipped |')
    out.append('| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |')
    g = {(r['scenario_id'], r['goal']): r for r in state_rows}
    for sid in dict.fromkeys(r['scenario_id'] for r in state_rows):
        t, d, rc = g[sid, 'tie'], g[sid, 'D+1-pair-margin'], g[sid, 'D-outside-automatic-recount']
        out.append(f"| {t['preference_contest']} | {t['volume_mix']} | {t['pair_pool']:,.1f} | {t['baseline_d_minus_r_net']:+,.1f} | {100 * t['solved_d_share_shift']:+.6f} | {100 * d['solved_d_share_shift']:+.6f} | {100 * rc['solved_d_share_shift']:+.6f} | {d['net_vote_change_needed']:,.1f} | {t['clipped_counties']} |")
    out.append('')
    out.append('<!-- table:contributors -->')
    out.append('| County | FIPS | 2020 Senate D−R net | Share of statewide 2020 net % | 2024 Senate D−R net | Share 2024 Senate % | 2024 president D−R net | Share 2024 pres % | 2022 governor D−R net | Share 2022 gov % |')
    out.append('| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: |')
    for n in sorted(NAMES, key=lambda n: cc[HEADLINE, n]['d_minus_r_net']):
        a, b, c, d = cc['2020-senate', n], cc['2024-senate', n], cc['2024-president', n], cc['2022-governor', n]
        out.append(f"| {n} | {COUNTY[n]['fips']} | {a['d_minus_r_net']:+,} | {100 * a['share_of_statewide_net']:.2f} | {b['d_minus_r_net']:+,} | {100 * b['share_of_statewide_net']:.2f} | {c['d_minus_r_net']:+,} | {100 * c['share_of_statewide_net']:.2f} | {d['d_minus_r_net']:+,} | {100 * d['share_of_statewide_net']:.2f} |")
    out.append('')
    out.append('<!-- table:selected -->')
    out.append('| Selected county | 2020 Senate pair votes | Other 22 counties fixed net | Tie: required D share % | D+1: required D share % | Status | Net if county were 100% D |')
    out.append('| --- | ---: | ---: | ---: | ---: | --- | ---: |')
    sel = {(r['county'], r['goal']): r for r in sel_rows if r['scenario_id'] == f'{HEADLINE}__native-pair'}
    for n in sorted(NAMES, key=lambda n: -sel[n, 'tie']['selected_pair_votes'])[:6] + ['Teton', 'Niobrara']:
        t, d = sel[n, 'tie'], sel[n, 'D+1-pair-margin']
        out.append(f"| {n} | {t['selected_pair_votes']:,.0f} | {t['fixed_other_counties_net']:+,.0f} | {100 * t['required_d_pair_share']:.2f} | {100 * d['required_d_pair_share']:.2f} | {t['status']} | {t['max_net_if_county_100pct_d']:+,.0f} |")
    out.append('')
    out.append('<!-- table:counties -->')
    out.append('| County | FIPS | 2020 Sen D % | 2020 Sen pp | 2024 Sen pp | 2024 Pres pp | 2022 Gov pp | 2022 House pp | 2024 House pp | 2020 pair weight % | 2020 net | +1 pp gain | Tie target D % | D+1 target D % | D+1 target net | +10% turnout net |')
    out.append('| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: |')
    for w in wide:
        n = w['county']
        out.append(f"| {n} | {w['county_fips']} | {100 * w['2020_senate_d_pair_share']:.2f} | {100 * w['2020_senate_d_minus_r_pair_margin']:+.2f} | {100 * w['2024_senate_d_minus_r_pair_margin']:+.2f} | {100 * w['2024_president_d_minus_r_pair_margin']:+.2f} | {100 * w['2022_governor_d_minus_r_pair_margin']:+.2f} | {100 * w['2022_house_d_minus_r_pair_margin']:+.2f} | {100 * w['2024_house_d_minus_r_pair_margin']:+.2f} | {100 * w['2020_senate_pair_weight']:.4f} | {w['2020_senate_d_minus_r_net']:+,} | {w['scn_2020senate_native_net_gain_per_plus1pp']:,.2f} | {100 * w['scn_2020senate_native_tie_target_d_share']:.4f} | {100 * w['scn_2020senate_native_dplus1_target_d_share']:.4f} | {w['scn_2020senate_native_dplus1_target_net']:+,.2f} | {w['scn_2020senate_net_change_per_plus10pct_pair_turnout']:+,.1f} |")
    return '\n'.join(out)


if __name__ == '__main__':
    main()
