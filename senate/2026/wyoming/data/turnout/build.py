# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5", "pypdf==6.21.1"]
# ///
"""Build Wyoming county turnout tables from retained official originals.

Run offline from any directory:
    uv run --with openpyxl==3.1.5 --with pypdf==6.21.1 python3 -I \
        senate/2026/wyoming/data/turnout/build.py

Inputs are only the files under originals/ listed in source-manifest.json.
Outputs (overwritten):
    ../county-turnout.csv              one row per county, ranked by 2024 U.S. Senate contest votes
    extracted-county-results.csv       every county cell used, with workbook member/sheet/cell
    precinct-reconciliation.csv        precinct-row sums vs county Total rows vs statewide summaries
    registration-snapshots.csv         county registration rows parsed from SOS registration PDFs
    rankings.json                      ordering, top-15 share and minimum-80% set for every measure
    build-checks.json                  every assertion evaluated during the build

Definitions (kept separate throughout):
    contest votes  = named candidates + aggregated "Write-Ins" for one contest (over/under votes excluded)
    ballots cast   = the SOS "Total Ballots Cast" summary for the election
    registration   = SOS registration snapshot (a stock of records at a date), never a turnout denominator here
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import logging
import re
import sys
import warnings
import zipfile
from decimal import ROUND_HALF_EVEN, Decimal, getcontext
from fractions import Fraction
from pathlib import Path

import openpyxl
import pypdf

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
logging.getLogger("pypdf").setLevel(logging.ERROR)
getcontext().prec = 50

HERE = Path(__file__).resolve().parent
ORIG = HERE / "originals"
OUT_CSV = HERE.parent / "county-turnout.csv"
RESEARCH_DATE = "2026-10-18"

CHECKS: list[dict] = []
ANNOTATIONS: list[dict] = []
DATE_COERCED: list[dict] = []


def check(name: str, ok: bool, detail=None, fatal: bool = True) -> bool:
    CHECKS.append({"check": name, "pass": bool(ok), "detail": detail})
    if not ok and fatal:
        raise SystemExit(f"CHECK FAILED: {name}: {detail}")
    return ok


def norm(value) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def unit_label(value) -> tuple[str, bool]:
    """Return (label, date_coerced). Some county sheets store precinct IDs such
    as 01-01 as Excel dates (2022-01-01); the label is rebuilt as MM-DD and
    flagged, and the raw cell value is retained in date_coerced_labels."""
    if isinstance(value, dt.datetime):
        return f"{value.month:02d}-{value.day:02d}", True
    return norm(value), False


def as_int(value, where: str) -> int:
    if isinstance(value, bool) or value is None:
        raise SystemExit(f"non-numeric/missing value at {where}: {value!r}")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    raise SystemExit(f"non-integer value at {where}: {value!r}")


def pct(num: int, den: int, places: int = 10) -> str:
    q = Decimal(1).scaleb(-places)
    return str((Decimal(num) * 100 / Decimal(den)).quantize(q, rounding=ROUND_HALF_EVEN))


# ---------------------------------------------------------------- geography
def load_counties() -> list[dict]:
    rows = []
    with open(ORIG / "st56_wy_cou2020.txt", newline="") as fh:
        for rec in csv.DictReader(fh, delimiter="|"):
            name = rec["COUNTYNAME"]
            assert name.endswith(" County"), name
            rows.append(
                {
                    "county": name[: -len(" County")],
                    "fips": rec["STATEFP"] + rec["COUNTYFP"],
                    "countyns": rec["COUNTYNS"],
                    "census_name": name,
                }
            )
    check("census: 23 Wyoming counties with unique 5-digit FIPS", len(rows) == 23 and len({r["fips"] for r in rows}) == 23
          and all(re.fullmatch(r"56\d{3}", r["fips"]) for r in rows), len(rows))
    return rows


COUNTIES = load_counties()
NAMES = [c["county"] for c in COUNTIES]
FIPS = {c["county"]: c["fips"] for c in COUNTIES}


# ---------------------------------------------------------------- workbooks
_WB_CACHE: dict = {}


def workbook(zip_name: str, member: str):
    key = (zip_name, member)
    if key not in _WB_CACHE:
        with zipfile.ZipFile(ORIG / zip_name) as zf:
            data = zf.read(member)
        # data_only=True reads each cell's cached value (2024 summary ballots
        # cells are cached results of formulas that pointed at a Google Sheet).
        _WB_CACHE[key] = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    return _WB_CACHE[key]


def find_title(ws, title: str, max_row: int = 8):
    for row in ws.iter_rows(min_row=1, max_row=max_row):
        for cell in row:
            if norm(cell.value) == title:
                return cell.row, cell.column
    raise SystemExit(f"title {title!r} not found in sheet {ws.title!r}")


def contest_block(ws, title: str, row_kind: str):
    """Parse one contest block: header row with contest title, labels below,
    then one row per county/precinct until a 'Total' row."""
    hrow, c0 = find_title(ws, title)
    c1 = ws.max_column
    for c in range(c0 + 1, ws.max_column + 1):
        if ws.cell(hrow, c).value not in (None, ""):
            c1 = c - 1
            break
    labels = []
    for c in range(c0, c1 + 1):
        lab = norm(ws.cell(hrow + 1, c).value)
        if not lab:
            c1 = c - 1
            break
        labels.append(lab)
    rows, total, r = [], None, hrow + 2
    blank_run = 0
    while r <= ws.max_row:
        key = ws.cell(r, 1).value
        vals = [ws.cell(r, c).value for c in range(c0, c0 + len(labels))]
        if key is None or norm(key) == "":
            if any(v is not None for v in vals):
                raise SystemExit(f"{ws.title}!row {r}: values without a row label")
            blank_run += 1
            if blank_run > 50:
                break
            r += 1
            continue
        blank_run = 0
        coords = [ws.cell(r, c).coordinate for c in range(c0, c0 + len(labels))]
        if all(v is None for v in vals):
            # layout annotation such as "Precincts Continue on Next Page": no counts, not a unit
            ANNOTATIONS.append({"sheet": ws.title, "row": r, "label": norm(key), "contest": title})
            r += 1
            continue
        ints = [as_int(v, f"{ws.title}!{coord}") for v, coord in zip(vals, coords)]
        if norm(key).lower() == "total":
            total = {"values": ints, "cells": coords, "row": r}
            break
        label, coerced = unit_label(key)
        if coerced:
            DATE_COERCED.append({"sheet": ws.title, "cell": ws.cell(r, 1).coordinate, "raw": key.isoformat(),
                                 "label_used": label, "contest": title})
        rows.append({"label": label, "values": ints, "cells": coords, "row": r, "date_coerced": coerced})
        r += 1
    if total is None:
        raise SystemExit(f"{ws.title}: no Total row for {title}")
    return {"labels": labels, "rows": rows, "total": total, "sheet": ws.title, "kind": row_kind}


def tbc_block(ws):
    hrow, col = find_title(ws, "Total Ballots Cast")
    nxt = [norm(ws.cell(hrow + 1, c).value) for c in range(1, ws.max_column + 1)]
    if "Total Ballots" in nxt:  # 2026 primary layout: party ballot columns under the title
        labels_row = hrow + 1
        cols = [c for c in range(2, ws.max_column + 1) if norm(ws.cell(labels_row, c).value)]
        labels = [norm(ws.cell(labels_row, c).value) for c in cols]
        start = labels_row + 1
    else:
        cols, labels, start = [col], ["Total Ballots Cast"], hrow + 1
        if norm(ws.cell(hrow + 1, 1).value) == "Precinct" or norm(ws.cell(hrow, 1).value) == "Precinct":
            start = hrow + 1
    rows, total = [], None
    for r in range(start, ws.max_row + 1):
        key, coerced = unit_label(ws.cell(r, 1).value)
        if coerced:
            DATE_COERCED.append({"sheet": ws.title, "cell": ws.cell(r, 1).coordinate,
                                 "raw": ws.cell(r, 1).value.isoformat(), "label_used": key, "contest": "Total Ballots Cast"})
        if not key or key == "Precinct":
            if any(ws.cell(r, c).value is not None for c in cols):
                raise SystemExit(f"{ws.title}!row {r}: ballots without label")
            continue
        coords = [ws.cell(r, c).coordinate for c in cols]
        if all(ws.cell(r, c).value is None for c in cols):
            ANNOTATIONS.append({"sheet": ws.title, "row": r, "label": key, "contest": "Total Ballots Cast"})
            continue
        vals = [as_int(ws.cell(r, c).value, f"{ws.title}!{ws.cell(r, c).coordinate}") for c in cols]
        if key.lower() == "total":
            total = {"values": vals, "cells": coords, "row": r}
            break
        rows.append({"label": key, "values": vals, "cells": coords, "row": r})
    if total is None:
        raise SystemExit(f"{ws.title}: no ballots Total row")
    return {"labels": labels, "rows": rows, "total": total, "sheet": ws.title}


# ---------------------------------------------------------------- election specs
ELECTIONS = {
    "2024": {
        "date": "2024-11-05",
        "zip": "2024_Wyoming_General_Results.zip",
        "summary": ("2024 General Results Summaries - OFFICIAL.xlsx", "Statewide Candidates"),
        "tbc": ("2024 General Results Summaries - OFFICIAL.xlsx", "Statewide Total Ballots Cast"),
        "pbp": ("2024 County General PbP - OFFICIAL.xlsx", None),
        "pbp_tbc": ("2024 County General PbP TBC - OFFICIAL.xlsx", None),
    },
    "2022": {
        "date": "2022-11-08",
        "zip": "2022_Wyoming_General_Results.zip",
        "summary": ("2022_General_Summaries_Results.xlsx", "Statewide Candidates"),
        "tbc": ("2022_General_Summaries_Results.xlsx", "Statewide Total Ballots Cast"),
        "pbp": ("2022_General_County_PbP.xlsx", None),
        "pbp_tbc": ("2022_General_County_PbP_TBC.xlsx", None),
    },
    "2020": {
        "date": "2020-11-03",
        "zip": "2020_Wyoming_General_Results.zip",
        "summary": ("2020_General_Statewide_Candidates_Summary.xlsx", "Sheet1"),
        "tbc": ("2020_General_Statewide_Total_Ballots_Cast_Summary.xlsx", "Sheet1"),
        "pbp": ("2020_{u}_County_General_PbP.xlsx", "Sheet1"),
        "pbp_tbc": ("2020_{u}_General_PbP_TBC.xlsx", "Sheet1"),
    },
}

CONTESTS = {
    "sen2024": {
        "year": "2024", "title": "United States Senator",
        "office": "U.S. Senate (Class I seat; not the 2026 Class II seat)",
        "labels": ["John Barrasso (R)", "Scott D Morrow (D)", "Write-Ins", "Overvotes", "Undervotes"],
        "fields": ["barrasso_r", "morrow_d", "write_ins", "overvotes", "undervotes"],
    },
    "pres2024": {
        "year": "2024", "title": "United States President",
        "office": "U.S. President",
        "labels": ["Donald J. Trump and JD Vance (R)", "Kamala D. Harris and Tim Walz (D)",
                   "Chase Oliver and Mike ter Maat (L)", "Write-Ins", "Overvotes", "Undervotes"],
        "fields": ["trump_vance_r", "harris_walz_d", "oliver_termaat_l", "write_ins", "overvotes", "undervotes"],
    },
    "gov2022": {
        "year": "2022", "title": "Governor",
        "office": "Governor (2022 midterm statewide contest; no 2022 U.S. Senate contest in Wyoming)",
        "labels": ["Mark Gordon (R)", "Theresa A. Livingston (D)", "Jared J. Baldes (L)", "Write-Ins",
                   "Overvotes", "Undervotes"],
        "fields": ["gordon_r", "livingston_d", "baldes_l", "write_ins", "overvotes", "undervotes"],
    },
    "sen2020": {
        "year": "2020", "title": "United States Senator",
        "office": "U.S. Senate (same Class II seat as 2026)",
        "labels": ["Cynthia M. Lummis (R)", "Merav Ben David (D)", "Write-Ins", "Overvotes", "Undervotes"],
        "fields": ["lummis_r", "ben_david_d", "write_ins", "overvotes", "undervotes"],
    },
}
NONCONTEST = {"overvotes", "undervotes"}

long_rows: list[dict] = []


def add_long(election, contest, county, label, category, value, member, sheet, cell):
    long_rows.append({
        "election_date": election, "contest": contest, "county": county,
        "county_fips": FIPS.get(county, ""), "source_column_label": label, "category": category,
        "value": value, "source_file": member, "sheet": sheet, "cell": cell,
    })


def category_of(field: str) -> str:
    if field == "write_ins":
        return "aggregated_write_ins"
    if field in NONCONTEST:
        return field[:-1]  # overvote / undervote
    return "named_candidate"


# ---------------------------------------------------------------- statewide summaries
county_data: dict[str, dict] = {n: {} for n in NAMES}
state_totals: dict[str, dict] = {}

for key, spec in CONTESTS.items():
    el = ELECTIONS[spec["year"]]
    member, sheet = el["summary"]
    ws = workbook(el["zip"], member)[sheet]
    blk = contest_block(ws, spec["title"], "county")
    check(f"{key}: candidate labels exactly as expected", blk["labels"] == spec["labels"], blk["labels"])
    labels_county = [r["label"] for r in blk["rows"]]
    check(f"{key}: 23 county rows in Census order", labels_county == NAMES, labels_county)
    sums = [sum(r["values"][i] for r in blk["rows"]) for i in range(len(spec["fields"]))]
    check(f"{key}: county rows sum to printed statewide Total row (every column)", sums == blk["total"]["values"],
          {"sum": sums, "total": blk["total"]["values"]})
    tot = dict(zip(spec["fields"], blk["total"]["values"]))
    tot["contest_votes"] = sum(v for f, v in tot.items() if f not in NONCONTEST)
    state_totals[key] = tot
    for r in blk["rows"]:
        d = dict(zip(spec["fields"], r["values"]))
        d["contest_votes"] = sum(v for f, v in d.items() if f not in NONCONTEST)
        county_data[r["label"]][key] = d
        for f, lab, v, cell in zip(spec["fields"], spec["labels"], r["values"], r["cells"]):
            add_long(el["date"], spec["office"], r["label"], lab, category_of(f), v, member, sheet, cell)

for year, el in ELECTIONS.items():
    member, sheet = el["tbc"]
    blk = tbc_block(workbook(el["zip"], member)[sheet])
    labels_county = [r["label"] for r in blk["rows"]]
    check(f"ballots{year}: 23 county rows in Census order", labels_county == NAMES, labels_county)
    s = sum(r["values"][0] for r in blk["rows"])
    check(f"ballots{year}: county ballots sum to printed Total", s == blk["total"]["values"][0], [s, blk["total"]["values"][0]])
    state_totals[f"ballots{year}"] = {"total_ballots_cast": blk["total"]["values"][0]}
    for r in blk["rows"]:
        county_data[r["label"]][f"ballots{year}"] = r["values"][0]
        add_long(el["date"], "Total Ballots Cast", r["label"], "Total Ballots Cast", "ballots_cast",
                 r["values"][0], member, sheet, r["cells"][0])

# 2026 primary ballots (context only; partisan primary ballots, not a general-election turnout)
P26_ZIP = "2026_Wyoming_Primary_Results.zip"
P26_MEMBER = "2026 Primary Results Summaries - OFFICIAL.xlsx"
blk = tbc_block(workbook(P26_ZIP, P26_MEMBER)["Statewide Total Ballots Cast"])
check("primary2026: ballot labels", blk["labels"] == ["Republican", "Democratic", "Non-Partisan", "Total Ballots"], blk["labels"])
check("primary2026: 23 county rows in Census order", [r["label"] for r in blk["rows"]] == NAMES)
for i, lab in enumerate(blk["labels"]):
    s = sum(r["values"][i] for r in blk["rows"])
    check(f"primary2026: county {lab} sum to printed Total", s == blk["total"]["values"][i], [s, blk["total"]["values"][i]])
for r in blk["rows"]:
    check(f"primary2026: {r['label']} party ballots sum to Total Ballots", sum(r["values"][:3]) == r["values"][3], r["values"])
    county_data[r["label"]]["primary2026"] = dict(zip(["republican", "democratic", "nonpartisan", "total"], r["values"]))
    for lab, v, cell in zip(blk["labels"], r["values"], r["cells"]):
        add_long("2026-08-18", "Primary Total Ballots Cast", r["label"], lab,
                 "primary_ballots_cast" if lab == "Total Ballots" else "primary_ballots_by_party_ballot",
                 v, P26_MEMBER, "Statewide Total Ballots Cast", cell)
state_totals["primary2026"] = dict(zip(["republican", "democratic", "nonpartisan", "total"], blk["total"]["values"]))

# ---------------------------------------------------------------- ballot identity (county level)
identity = []
for key, spec in CONTESTS.items():
    year = spec["year"]
    for n in NAMES:
        d = county_data[n][key]
        marks = d["contest_votes"] + d["overvotes"] + d["undervotes"]
        identity.append({"contest": key, "county": n, "contest_plus_over_under": marks,
                         "ballots_cast": county_data[n][f"ballots{year}"],
                         "difference": marks - county_data[n][f"ballots{year}"]})
bad = [r for r in identity if r["difference"] != 0]
check("county identity: contest votes + overvotes + undervotes == total ballots cast (all 4 contests x 23 counties)",
      not bad, bad, fatal=False)

# ---------------------------------------------------------------- precinct reconciliation
precinct_rows = []
precinct_identity_issues = []
for key, spec in CONTESTS.items():
    year = spec["year"]
    el = ELECTIONS[year]
    for n in NAMES:
        u = n.replace(" ", "_")
        pm, ps = el["pbp"]
        tm, ts = el["pbp_tbc"]
        pm, tm = pm.format(u=u), tm.format(u=u)
        ws = workbook(el["zip"], pm)[ps or n]
        blk = contest_block(ws, spec["title"], "precinct")
        check(f"{key} precinct {n}: labels", blk["labels"] == spec["labels"], blk["labels"])
        psum = [sum(r["values"][i] for r in blk["rows"]) for i in range(len(spec["fields"]))]
        tws = workbook(el["zip"], tm)[ts or n]
        tb = tbc_block(tws)
        tsum = sum(r["values"][0] for r in tb["rows"])
        summary_vals = [county_data[n][key][f] for f in spec["fields"]]
        ok_rows = psum == blk["total"]["values"] == summary_vals
        ok_tbc = tsum == tb["total"]["values"][0] == county_data[n][f"ballots{year}"]
        check(f"{key} precinct {n}: precinct sums == sheet Total == statewide county row", ok_rows,
              {"precinct_sum": psum, "sheet_total": blk["total"]["values"], "summary": summary_vals}, fatal=False)
        check(f"ballots{year} precinct {n}: precinct ballots sum == sheet Total == statewide county ballots", ok_tbc,
              {"precinct_sum": tsum, "sheet_total": tb["total"]["values"][0], "summary": county_data[n][f"ballots{year}"]},
              fatal=False)
        # precinct-level identity, matched by exact native precinct label
        tmap = {}
        for r in tb["rows"]:
            tmap.setdefault(r["label"], []).append(r["values"][0])
        cmap = {}
        for r in blk["rows"]:
            cmap.setdefault(r["label"], []).append(sum(r["values"]))
        dup = sorted({k for k, v in tmap.items() if len(v) > 1} | {k for k, v in cmap.items() if len(v) > 1})
        only_contest = sorted(set(cmap) - set(tmap))
        only_tbc = sorted(set(tmap) - set(cmap))
        diffs = []
        for lab in sorted(set(cmap) & set(tmap)):
            if len(cmap[lab]) == 1 and len(tmap[lab]) == 1 and cmap[lab][0] != tmap[lab][0]:
                diffs.append({"precinct": lab, "contest_plus_over_under": cmap[lab][0], "ballots_cast": tmap[lab][0]})
        if dup or only_contest or only_tbc or diffs:
            precinct_identity_issues.append({"contest": key, "county": n, "duplicate_labels": dup,
                                             "labels_only_in_contest_sheet": only_contest,
                                             "labels_only_in_ballots_sheet": only_tbc, "identity_differences": diffs})
        precinct_rows.append({
            "election_date": el["date"], "contest": key, "county": n, "county_fips": FIPS[n],
            "contest_workbook": pm, "contest_sheet": ws.title, "contest_precinct_rows": len(blk["rows"]),
            "contest_date_coerced_labels": sum(1 for r in blk["rows"] if r["date_coerced"]),
            "contest_precinct_sum_contest_votes": sum(v for f, v in zip(spec["fields"], psum) if f not in NONCONTEST),
            "contest_sheet_total_contest_votes": sum(v for f, v in zip(spec["fields"], blk["total"]["values"]) if f not in NONCONTEST),
            "statewide_summary_contest_votes": county_data[n][key]["contest_votes"],
            "ballots_workbook": tm, "ballots_sheet": tws.title, "ballots_precinct_rows": len(tb["rows"]),
            "ballots_precinct_sum": tsum, "ballots_sheet_total": tb["total"]["values"][0],
            "statewide_summary_ballots": county_data[n][f"ballots{year}"],
            "precinct_label_or_identity_issue": "yes" if (dup or only_contest or only_tbc or diffs) else "no",
            "contest_rows_reconciled": "yes" if ok_rows else "no",
            "ballots_rows_reconciled": "yes" if ok_tbc else "no",
        })

# 2026 primary precinct ballots reconciliation
P26_PBP_TBC = "2026 Primary County PbP Total Ballots Cast - OFFICIAL.xlsx"
for n in NAMES:
    tb = tbc_block(workbook(P26_ZIP, P26_PBP_TBC)[n])
    sums = [sum(r["values"][i] for r in tb["rows"]) for i in range(len(tb["labels"]))]
    target = [county_data[n]["primary2026"][k] for k in ("republican", "democratic", "nonpartisan", "total")]
    check(f"primary2026 precinct {n}: precinct ballots sum == sheet Total == statewide county row",
          tb["labels"] == ["Republican", "Democratic", "Non-Partisan", "Total Ballots"] and sums == tb["total"]["values"] == target,
          {"labels": tb["labels"], "sum": sums, "total": tb["total"]["values"], "summary": target}, fatal=False)
    precinct_rows.append({
        "election_date": "2026-08-18", "contest": "primary2026_ballots", "county": n, "county_fips": FIPS[n],
        "contest_workbook": "", "contest_sheet": "", "contest_precinct_rows": "", "contest_date_coerced_labels": "",
        "contest_precinct_sum_contest_votes": "", "contest_sheet_total_contest_votes": "", "statewide_summary_contest_votes": "",
        "ballots_workbook": P26_PBP_TBC, "ballots_sheet": n, "ballots_precinct_rows": len(tb["rows"]),
        "ballots_precinct_sum": sums[3], "ballots_sheet_total": tb["total"]["values"][3],
        "statewide_summary_ballots": target[3], "precinct_label_or_identity_issue": "not tested",
        "contest_rows_reconciled": "",
        "ballots_rows_reconciled": "yes" if sums == tb["total"]["values"] == target else "no",
    })


# ---------------------------------------------------------------- registration and history PDFs
def pdf_lines(name: str):
    reader = pypdf.PdfReader(ORIG / name)
    return [(i + 1, (p.extract_text(extraction_mode="layout") or "").splitlines()) for i, p in enumerate(reader.pages)]


NUM = r"\d{1,3}(?:,\d{3})*"


def registration_page(name: str, date_label: str, require_heading: str | None):
    for page_no, lines in pdf_lines(name):
        text = "\n".join(lines)
        if date_label not in text or (require_heading and require_heading not in text):
            continue
        header = next((l for l in lines if l.strip().startswith("County") and "TOTAL" in l), None)
        cols = header.split()[1:] if header else []
        cols = [c for c in cols if c not in ("Labels",)]
        cols = ["No Labels" if c == "No" else c for c in cols]
        rows = {}
        for l in lines:
            m = re.match(r"^\s*([A-Z][A-Z ]+?)\s{2,}(" + NUM + r"(?:\s+" + NUM + r")*)\s*$", l)
            if m:
                rows[m.group(1).strip()] = [int(x.replace(",", "")) for x in m.group(2).split()]
        return page_no, cols, rows, text
    raise SystemExit(f"{name}: registration page {date_label} not found")


REG_SPECS = {
    "reg2024": ("2024VR_stats.pdf", "11/5/2024", "General Election Day"),
    "reg2022": ("2022VR_stats.pdf", "11/8/2022", "General Election Day"),
    "reg2020": ("2020VR_stats.pdf", "11/3/2020", "General Election Day"),
    "reg2026_10_01": ("26OctVR_Stats.pdf", "10/1/2026", None),
}
reg_rows = []
for key, (fname, dlabel, heading) in REG_SPECS.items():
    page_no, cols, rows, _ = registration_page(fname, dlabel, heading)
    check(f"{key}: header columns parsed and end with TOTAL", bool(cols) and cols[-1] == "TOTAL", cols)
    upper = [n.upper() for n in NAMES]
    check(f"{key}: 23 county rows + TOTAL", set(rows) == set(upper) | {"TOTAL"}, sorted(rows))
    for n in NAMES:
        vals = rows[n.upper()]
        check(f"{key}: {n} row width", len(vals) == len(cols), vals)
        check(f"{key}: {n} party columns sum to TOTAL", sum(vals[:-1]) == vals[-1], vals)
        county_data[n][key] = vals[-1]
        rec = {"snapshot": key, "page_date_label": dlabel, "page_heading": heading or "", "county": n,
               "county_fips": FIPS[n], "source_file": fname, "pdf_page": page_no}
        rec.update({c: v for c, v in zip(cols, vals)})
        reg_rows.append(rec)
    tot = rows["TOTAL"]
    colsum = [sum(rows[n.upper()][i] for n in NAMES) for i in range(len(cols))]
    check(f"{key}: county columns sum to printed TOTAL row", colsum == tot, {"sum": colsum, "printed": tot})
    state_totals[key] = {"total_registered": tot[-1], "pdf_page": page_no, "columns": cols}

# Voter profile history: statewide registered (close of business day before) and turnout (ballots)
vp_text_plain = "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(ORIG / "VoterProfile.pdf").pages)
profile = {}
for line in vp_text_plain.splitlines():
    m = re.match(r"^(20\d\d|19\d\d)\s+(.*)$", line.strip())
    if not m:
        continue
    toks = re.findall(NUM + r"(?:\.\s?\d+)?%?", m.group(2))
    nums = [t for t in toks if not t.endswith("%")]
    profile[m.group(1)] = {"tokens": toks, "counts": [int(t.replace(",", "")) for t in nums]}
for year in ("2024", "2022", "2020"):
    counts = profile[year]["counts"]  # VAP, primary reg, primary turnout, general reg, general turnout
    check(f"voter profile {year}: five counts parsed", len(counts) == 5, profile[year])
    state_totals[f"profile{year}"] = {"voting_age_population_estimate": counts[0], "primary_registered": counts[1],
                                      "primary_turnout": counts[2], "general_registered": counts[3],
                                      "general_turnout": counts[4]}
    check(f"voter profile {year}: general turnout == statewide Total Ballots Cast",
          counts[4] == state_totals[f"ballots{year}"]["total_ballots_cast"], [counts[4], state_totals[f"ballots{year}"]])
    check(f"voter profile {year}: general registered == General Election Day registration TOTAL",
          counts[3] == state_totals[f"reg{year}"]["total_registered"], [counts[3], state_totals[f"reg{year}"]])
p26 = profile["2026"]["counts"]
state_totals["profile2026"] = {"voting_age_population_estimate": p26[0], "primary_registered": p26[1], "primary_turnout": p26[2]}
check("voter profile 2026: primary turnout == 2026 primary Total Ballots", p26[2] == state_totals["primary2026"]["total"], p26)
check("voter profile footnote: registration excludes election-day registrations",
      "does not" in vp_text_plain and "election day registrations" in vp_text_plain.replace("\n", " "), None)

# ---------------------------------------------------------------- rankings
MEASURES = {
    "sen2024_contest_votes": ("2024 U.S. Senate contest votes (Barrasso, Morrow, aggregated write-ins)",
                              lambda n: county_data[n]["sen2024"]["contest_votes"]),
    "pres2024_contest_votes": ("2024 U.S. President contest votes (three tickets plus aggregated write-ins)",
                               lambda n: county_data[n]["pres2024"]["contest_votes"]),
    "ballots2024": ("2024 general total ballots cast", lambda n: county_data[n]["ballots2024"]),
    "gov2022_contest_votes": ("2022 Governor contest votes (Gordon, Livingston, Baldes, aggregated write-ins)",
                              lambda n: county_data[n]["gov2022"]["contest_votes"]),
    "ballots2022": ("2022 general total ballots cast", lambda n: county_data[n]["ballots2022"]),
    "sen2020_contest_votes": ("2020 U.S. Senate contest votes (Lummis, Ben David, aggregated write-ins)",
                              lambda n: county_data[n]["sen2020"]["contest_votes"]),
    "ballots2020": ("2020 general total ballots cast", lambda n: county_data[n]["ballots2020"]),
    "primary2026_ballots": ("2026 primary total ballots cast (all party ballots)",
                            lambda n: county_data[n]["primary2026"]["total"]),
    # sensitivity: named candidates only (aggregated write-ins excluded)
    "sen2024_named_candidate_votes": ("SENSITIVITY: 2024 U.S. Senate named-candidate votes only (Barrasso + Morrow)",
                                      lambda n: county_data[n]["sen2024"]["barrasso_r"] + county_data[n]["sen2024"]["morrow_d"]),
    "gov2022_named_candidate_votes": ("SENSITIVITY: 2022 Governor named-candidate votes only (Gordon + Livingston + Baldes)",
                                      lambda n: sum(county_data[n]["gov2022"][f] for f in ("gordon_r", "livingston_d", "baldes_l"))),
}
OFFICIAL_TOTAL = {
    "sen2024_contest_votes": state_totals["sen2024"]["contest_votes"],
    "pres2024_contest_votes": state_totals["pres2024"]["contest_votes"],
    "ballots2024": state_totals["ballots2024"]["total_ballots_cast"],
    "gov2022_contest_votes": state_totals["gov2022"]["contest_votes"],
    "ballots2022": state_totals["ballots2022"]["total_ballots_cast"],
    "sen2020_contest_votes": state_totals["sen2020"]["contest_votes"],
    "ballots2020": state_totals["ballots2020"]["total_ballots_cast"],
    "primary2026_ballots": state_totals["primary2026"]["total"],
    "sen2024_named_candidate_votes": state_totals["sen2024"]["barrasso_r"] + state_totals["sen2024"]["morrow_d"],
    "gov2022_named_candidate_votes": sum(state_totals["gov2022"][f] for f in ("gordon_r", "livingston_d", "baldes_l")),
}


def rank_measure(key: str):
    label, fn = MEASURES[key]
    vals = {n: fn(n) for n in NAMES}
    total = sum(vals.values())
    check(f"{key}: county values sum to official statewide total", total == OFFICIAL_TOTAL[key], [total, OFFICIAL_TOTAL[key]])
    order = sorted(NAMES, key=lambda n: (-vals[n], FIPS[n]))
    cum, rows, k80 = 0, [], None
    for i, n in enumerate(order, 1):
        cum += vals[n]
        rows.append({"rank": i, "county": n, "county_fips": FIPS[n], "value": vals[n], "cumulative": cum,
                     "share_pct": pct(vals[n], total), "cumulative_pct": pct(cum, total)})
        if k80 is None and 5 * cum >= 4 * total:
            k80 = i
    before = rows[k80 - 2]["cumulative"] if k80 > 1 else 0
    at = rows[k80 - 1]["cumulative"]
    check(f"{key}: minimum-80 integer proof", 5 * before < 4 * total <= 5 * at, [before, total, at])
    top15 = rows[14]["cumulative"]
    ties = sorted({vals[n] for n in NAMES if list(vals.values()).count(vals[n]) > 1})
    return {
        "measure": key, "definition": label, "statewide_total": total, "official_statewide_total": OFFICIAL_TOTAL[key],
        "top15": order[:15], "top15_votes": top15, "top15_share": f"{top15}/{total}", "top15_pct": pct(top15, total),
        "minimum80_k": k80, "minimum80_set": order[:k80], "minimum80_votes": at, "minimum80_pct": pct(at, total),
        "preceding_k_minus_1_votes": before, "preceding_k_minus_1_pct": pct(before, total),
        "crossing_county": order[k80 - 1], "crossing_county_fips": FIPS[order[k80 - 1]],
        "integer_proof": f"5*{before}={5 * before} < 4*{total}={4 * total} <= 5*{at}={5 * at}",
        "remainder_k": 23 - k80, "remainder_votes": total - at, "remainder_pct": pct(total - at, total),
        "tied_values": ties, "tie_rule": "descending value, then ascending county FIPS",
        "rows": rows,
    }


rankings = {k: rank_measure(k) for k in MEASURES}
base = rankings["sen2024_contest_votes"]
check("primary measure: top 15 contained in the minimum-80 set or vice versa",
      set(base["minimum80_set"]) <= set(base["top15"]) or set(base["top15"]) <= set(base["minimum80_set"]), None)
priority = [n for n in base["rows"] if n["county"] in set(base["top15"]) | set(base["minimum80_set"])]
pv = sum(r["value"] for r in priority)
rankings["_priority_union"] = {
    "definition": "union of the top 15 and the minimum set reaching 80% of 2024 U.S. Senate contest votes",
    "counties": [r["county"] for r in priority], "fips": [r["county_fips"] for r in priority],
    "k": len(priority), "votes": pv, "share": f"{pv}/{base['statewide_total']}",
    "pct": pct(pv, base["statewide_total"]),
}
for key in MEASURES:
    if key == "sen2024_contest_votes":
        continue
    r = rankings[key]
    r["minimum80_set_equal_to_sen2024"] = set(r["minimum80_set"]) == set(base["minimum80_set"])
    r["minimum80_added_vs_sen2024"] = sorted(set(r["minimum80_set"]) - set(base["minimum80_set"]))
    r["minimum80_missing_vs_sen2024"] = sorted(set(base["minimum80_set"]) - set(r["minimum80_set"]))
    r["top15_set_equal_to_sen2024"] = set(r["top15"]) == set(base["top15"])
    r["top15_added_vs_sen2024"] = sorted(set(r["top15"]) - set(base["top15"]))
    r["top15_missing_vs_sen2024"] = sorted(set(base["top15"]) - set(r["top15"]))
    r["all_county_order_equal_to_sen2024"] = [x["county"] for x in r["rows"]] == [x["county"] for x in base["rows"]]
    r["priority_union_votes_in_this_measure"] = sum(x["value"] for x in r["rows"] if x["county"] in rankings["_priority_union"]["counties"])
    r["priority_union_pct_in_this_measure"] = pct(r["priority_union_votes_in_this_measure"], r["statewide_total"])

# ---------------------------------------------------------------- outputs
rank_of = {k: {x["county"]: x["rank"] for x in v["rows"]} for k, v in rankings.items() if not k.startswith("_")}
in80 = {k: set(v["minimum80_set"]) for k, v in rankings.items() if not k.startswith("_")}
in15 = {k: set(v["top15"]) for k, v in rankings.items() if not k.startswith("_")}
share = {k: {x["county"]: x["share_pct"] for x in v["rows"]} for k, v in rankings.items() if not k.startswith("_")}
yn = lambda b: "yes" if b else "no"  # noqa: E731

fields = [
    "rank_sen2024", "county", "county_fips", "census_countyns",
    "sen2024_barrasso_r", "sen2024_morrow_d", "sen2024_write_ins", "sen2024_contest_votes",
    "sen2024_overvotes", "sen2024_undervotes", "sen2024_share_pct", "sen2024_cumulative_contest_votes",
    "sen2024_cumulative_pct", "in_top15_sen2024", "in_minimum80_sen2024", "in_priority_union",
    "pres2024_contest_votes", "pres2024_overvotes", "pres2024_undervotes", "rank_pres2024", "in_minimum80_pres2024",
    "ballots2024_total_ballots_cast", "rank_ballots2024", "ballots2024_share_pct", "in_minimum80_ballots2024",
    "reg2024_11_05_general_election_day_total",
    "gov2022_gordon_r", "gov2022_livingston_d", "gov2022_baldes_l", "gov2022_write_ins", "gov2022_contest_votes",
    "gov2022_overvotes", "gov2022_undervotes", "rank_gov2022", "gov2022_share_pct", "in_top15_gov2022",
    "in_minimum80_gov2022",
    "ballots2022_total_ballots_cast", "rank_ballots2022", "ballots2022_share_pct", "in_minimum80_ballots2022",
    "reg2022_11_08_general_election_day_total",
    "sen2020_lummis_r", "sen2020_ben_david_d", "sen2020_write_ins", "sen2020_contest_votes",
    "sen2020_overvotes", "sen2020_undervotes", "rank_sen2020", "sen2020_share_pct", "in_minimum80_sen2020",
    "ballots2020_total_ballots_cast", "rank_ballots2020", "in_minimum80_ballots2020",
    "reg2020_11_03_general_election_day_total",
    "primary2026_republican_ballots", "primary2026_democratic_ballots", "primary2026_nonpartisan_ballots",
    "primary2026_total_ballots_cast", "rank_primary2026_ballots",
    "reg2026_10_01_total",
    "source_status", "research_date",
]
out_rows = []
cum_by = {x["county"]: x for x in base["rows"]}
ns = {c["county"]: c["countyns"] for c in COUNTIES}
for x in base["rows"]:
    n = x["county"]
    d = county_data[n]
    s24, p24, g22, s20, p26 = d["sen2024"], d["pres2024"], d["gov2022"], d["sen2020"], d["primary2026"]
    out_rows.append({
        "rank_sen2024": x["rank"], "county": n, "county_fips": FIPS[n], "census_countyns": ns[n],
        "sen2024_barrasso_r": s24["barrasso_r"], "sen2024_morrow_d": s24["morrow_d"], "sen2024_write_ins": s24["write_ins"],
        "sen2024_contest_votes": s24["contest_votes"], "sen2024_overvotes": s24["overvotes"],
        "sen2024_undervotes": s24["undervotes"], "sen2024_share_pct": x["share_pct"],
        "sen2024_cumulative_contest_votes": x["cumulative"], "sen2024_cumulative_pct": x["cumulative_pct"],
        "in_top15_sen2024": yn(n in in15["sen2024_contest_votes"]),
        "in_minimum80_sen2024": yn(n in in80["sen2024_contest_votes"]),
        "in_priority_union": yn(n in rankings["_priority_union"]["counties"]),
        "pres2024_contest_votes": p24["contest_votes"], "pres2024_overvotes": p24["overvotes"],
        "pres2024_undervotes": p24["undervotes"], "rank_pres2024": rank_of["pres2024_contest_votes"][n],
        "in_minimum80_pres2024": yn(n in in80["pres2024_contest_votes"]),
        "ballots2024_total_ballots_cast": d["ballots2024"], "rank_ballots2024": rank_of["ballots2024"][n],
        "ballots2024_share_pct": share["ballots2024"][n], "in_minimum80_ballots2024": yn(n in in80["ballots2024"]),
        "reg2024_11_05_general_election_day_total": d["reg2024"],
        "gov2022_gordon_r": g22["gordon_r"], "gov2022_livingston_d": g22["livingston_d"], "gov2022_baldes_l": g22["baldes_l"],
        "gov2022_write_ins": g22["write_ins"], "gov2022_contest_votes": g22["contest_votes"],
        "gov2022_overvotes": g22["overvotes"], "gov2022_undervotes": g22["undervotes"],
        "rank_gov2022": rank_of["gov2022_contest_votes"][n], "gov2022_share_pct": share["gov2022_contest_votes"][n],
        "in_top15_gov2022": yn(n in in15["gov2022_contest_votes"]),
        "in_minimum80_gov2022": yn(n in in80["gov2022_contest_votes"]),
        "ballots2022_total_ballots_cast": d["ballots2022"], "rank_ballots2022": rank_of["ballots2022"][n],
        "ballots2022_share_pct": share["ballots2022"][n], "in_minimum80_ballots2022": yn(n in in80["ballots2022"]),
        "reg2022_11_08_general_election_day_total": d["reg2022"],
        "sen2020_lummis_r": s20["lummis_r"], "sen2020_ben_david_d": s20["ben_david_d"], "sen2020_write_ins": s20["write_ins"],
        "sen2020_contest_votes": s20["contest_votes"], "sen2020_overvotes": s20["overvotes"],
        "sen2020_undervotes": s20["undervotes"], "rank_sen2020": rank_of["sen2020_contest_votes"][n],
        "sen2020_share_pct": share["sen2020_contest_votes"][n], "in_minimum80_sen2020": yn(n in in80["sen2020_contest_votes"]),
        "ballots2020_total_ballots_cast": d["ballots2020"], "rank_ballots2020": rank_of["ballots2020"][n],
        "in_minimum80_ballots2020": yn(n in in80["ballots2020"]),
        "reg2020_11_03_general_election_day_total": d["reg2020"],
        "primary2026_republican_ballots": p26["republican"], "primary2026_democratic_ballots": p26["democratic"],
        "primary2026_nonpartisan_ballots": p26["nonpartisan"], "primary2026_total_ballots_cast": p26["total"],
        "rank_primary2026_ballots": rank_of["primary2026_ballots"][n],
        "reg2026_10_01_total": d["reg2026_10_01"],
        "source_status": "WY SOS official general results (2024 certified by State Canvassing Board 2024-11-13; 2022/2020 pages titled Official); 2026 primary official; registration SOS monthly statistics",
        "research_date": RESEARCH_DATE,
    })


def write_csv(path: Path, rows: list[dict], header: list[str] | None = None):
    header = header or list(rows[0].keys())
    with open(path, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=header, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


write_csv(OUT_CSV, out_rows, fields)


def disp(p: str, places: int = 4) -> str:
    return str(Decimal(p).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_EVEN))


table = [
    "| Rank | County | FIPS | 2024 Senate contest votes | Share | Cumulative | 2024 ballots cast | "
    "2022 Governor contest votes (rank) | 2022 share | 2020 Senate contest votes (rank) | Priority |",
    "| ---: | --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | --- |",
]
for r in out_rows:
    tag = ("top 15, 80% set" if r["in_minimum80_sen2024"] == "yes" else "top 15") if r["in_top15_sen2024"] == "yes" else ""
    table.append(
        f"| {r['rank_sen2024']} | {r['county']} | {r['county_fips']} | {r['sen2024_contest_votes']:,} | "
        f"{disp(r['sen2024_share_pct'])}% | {disp(r['sen2024_cumulative_pct'])}% | {r['ballots2024_total_ballots_cast']:,} | "
        f"{r['gov2022_contest_votes']:,} ({r['rank_gov2022']}) | {disp(r['gov2022_share_pct'])}% | "
        f"{r['sen2020_contest_votes']:,} ({r['rank_sen2020']}) | {tag} |"
    )
tot = state_totals
table.append(
    f"| | **Statewide** | 56 | **{tot['sen2024']['contest_votes']:,}** | 100% | | "
    f"**{tot['ballots2024']['total_ballots_cast']:,}** | **{tot['gov2022']['contest_votes']:,}** | 100% | "
    f"**{tot['sen2020']['contest_votes']:,}** | |"
)
(HERE / "ranking-table.md").write_text(
    "<!-- Generated by build.py from ../county-turnout.csv; display rounding only. -->\n" + "\n".join(table) + "\n"
)
write_csv(HERE / "extracted-county-results.csv", long_rows)
write_csv(HERE / "precinct-reconciliation.csv", precinct_rows)
reg_header = ["snapshot", "page_date_label", "page_heading", "county", "county_fips", "Constitution", "Democratic",
              "Libertarian", "No Labels", "Republican", "Unaffiliated", "Other*", "TOTAL", "source_file", "pdf_page"]
write_csv(HERE / "registration-snapshots.csv", [{k: r.get(k, "") for k in reg_header} for r in reg_rows], reg_header)

meta = {
    "research_date": RESEARCH_DATE,
    "state_totals": state_totals,
    "county_identity_differences": bad,
    "precinct_identity_issues": precinct_identity_issues,
    "layout_annotation_rows_skipped": ANNOTATIONS,
    "date_coerced_precinct_labels": DATE_COERCED,
    "rankings": rankings,
}
(HERE / "rankings.json").write_text(json.dumps(meta, indent=2) + "\n")
(HERE / "build-checks.json").write_text(json.dumps({
    "generated_by": "build.py", "research_date": RESEARCH_DATE,
    "checks_total": len(CHECKS), "checks_failed": sum(1 for c in CHECKS if not c["pass"]),
    "checks": CHECKS,
}, indent=2) + "\n")

print(f"checks: {len(CHECKS)} evaluated, {sum(1 for c in CHECKS if not c['pass'])} non-fatal failures")
print(f"county identity differences: {len(bad)}; precinct label/identity issue counties: {len(precinct_identity_issues)}")
for k, r in rankings.items():
    if k.startswith("_"):
        print(k, r["k"], r["pct"], r["counties"])
        continue
    print(f"{k}: total {r['statewide_total']}, top15 {r['top15_pct']}%, min80 k={r['minimum80_k']} "
          f"{r['minimum80_pct']}% crossing {r['crossing_county']}")
sys.exit(0)
